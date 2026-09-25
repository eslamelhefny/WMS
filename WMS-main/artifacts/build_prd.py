from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'artifacts' / 'Andalusia_Academy_PRD.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin, sec.bottom_margin = Inches(.73), Inches(.65)
sec.left_margin = sec.right_margin = Inches(.73)
sec.header_distance = sec.footer_distance = Inches(.3)
PRIMARY, DARK, SOFT, INK = 'C17A62', '8F4F3B', 'F3E4DF', '55372F'

for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor.from_string('000000')
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
normal = doc.styles['Normal']
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.1
for name, size, before, after in [('Title',27,6,12),('Subtitle',14,0,12),('Heading 1',20,0,14),('Heading 2',12,10,5),('Heading 3',11,7,4)]:
    st=doc.styles[name];st.font.size=Pt(size)
    st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after)
    st.paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].font.color.rgb=RGBColor.from_string('7D6A63')

header=sec.header.paragraphs[0]
header.text='ANDALUSIA ACADEMY    /    PRODUCT REQUIREMENTS'
header.runs[0].font.size=Pt(8)
header.runs[0].font.bold=True
footer=sec.footer.paragraphs[0]
footer.paragraph_format.space_before=Pt(5)
footer.add_run('WMS  |  Version 1.0  |  8 September 2026').font.size=Pt(8)
footer.add_run(' '*8+'Page ').font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for run in footer.runs:run.font.color.rgb=RGBColor.from_string(DARK)

def p(text='',bold=False,style=None):
    para=doc.add_paragraph(style=style)
    run=para.add_run(text);run.bold=bold
    return para
def h(text):doc.add_heading(text,2)
def bullet(text):
    para=p(text,style='List Bullet');para.paragraph_format.space_after=Pt(5)
def page(number,title):
    doc.add_page_break()
    e=p(f'{number:02d}   ANDALUSIA ACADEMY')
    e.runs[0].font.color.rgb=RGBColor.from_string(DARK);e.runs[0].font.size=Pt(9);e.runs[0].bold=True
    doc.add_heading(title,1)
def req(id,title,text):
    para=doc.add_paragraph();para.add_run(f'{id}  {title}  ').bold=True;para.add_run(text)
def table(headers, rows, widths, size=9.5):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'5');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    pr.append(borders)
    for index, vals in enumerate([headers]+rows):
        row=t.rows[0] if index==0 else t.add_row()
        trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
        if index==0:
            rep=OxmlElement('w:tblHeader');trPr.append(rep)
        for j,(cell,value,width) in enumerate(zip(row.cells,vals,widths)):
            cell.width=Inches(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),DARK if index==0 else ('FFF9F6' if index%2 else 'FFFFFF'));cp.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                m=OxmlElement('w:'+side);m.set(qn('w:w'),'95');m.set(qn('w:type'),'dxa');margins.append(m)
            cp.append(margins)
            para=cell.paragraphs[0];para.paragraph_format.space_after=Pt(1);para.paragraph_format.line_spacing=1.05
            r=para.add_run(str(value));r.font.size=Pt(size);r.bold=index==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if index==0 else INK)
    p().paragraph_format.space_after=Pt(1)
    return t

logo=p();logo.alignment=WD_ALIGN_PARAGRAPH.CENTER
pic=logo.add_run().add_picture(str(ROOT/'static/logo.png'),width=Inches(1.8))
pic._inline.docPr.set('descr','Original Andalusia Academy bilingual logo in terracotta')
doc.add_paragraph('Andalusia Academy\nWork Management System',style='Title')
doc.add_paragraph('Product Requirements Document',style='Subtitle')
p('Version 1.0   |   8 September 2026   |   Draft for stakeholder review',True)
h('Purpose and product direction')
p('This document defines the requirements for Andalusia Academy’s Work Management System, including the new Plan dashboard and Visits networking database. It gives Academy leadership, operations teams, designers, developers, and testers a shared baseline for product scope, expected behavior, and acceptance.')
p('The product brings program delivery, position-owned work, approvals, and relationship records into one workspace. Plan makes active and upcoming work visible. Visits preserves contact details, visit outcomes, and next steps so relationships can be followed across programs and activities.')
h('Business outcomes')
bullet('Give managers a clear view of program readiness, work in progress, upcoming commitments, and overdue follow-ups.')
bullet('Keep reusable contact records connected to visit history, program context, and operational activity.')
bullet('Preserve the eight-stage Academy lifecycle and its evidence and approval requirements.')
h('Document basis')
p('This PRD is grounded in the local application and the requested Plan and Visits enhancement. “Baseline” describes implemented behavior. “Production requirement” identifies work or a decision still needed before shared organizational deployment. This document does not record business approval or a production release.')
p('Sections 2 to 4 define the system scope. Sections 5 to 7 specify Plan, Visits, and their data. Sections 8 to 10 cover quality, acceptance, and release decisions.',style='Caption')

page(2,'Users and scope')
p('The system organizes responsibility by Academy Position. The active Position and additional Positions influence navigation, assignments, and stage actions. Networking records currently form a shared directory within the local application.')
table(['User group','Primary responsibility'],[
('Academy Manager','Review portfolio health, allocate work, make key decisions, and complete lifecycle stages after validation.'),
('Program executives and supervisors','Create and coordinate programs where their configured Position allows; monitor preparation and delivery.'),
('Curriculum and instructor teams','Maintain curriculum, quality checks, instructor selection, contracts, and onboarding evidence.'),
('Operations specialists','Own instructor coordination, logistics and place, attendees, or sponsorship activities and their interactions.'),
('Commercial and partnerships teams','Maintain commercial readiness and relationships with clients, partners, and sponsors.'),
('Finance PMB and Treasury','Maintain distinct financial-result and settlement responsibilities.'),
('All registered local users','Manage their own position-based work and use the currently shared Plan and Visits workspaces.')
],[2.05,4.75])
h('Included in this release baseline')
p('Program lifecycle management; the Manager Dashboard; tasks and team views; calendar planning; instructors; operations; finance; delivery scheduling; documents and reports; history; profile and administration; and the new Plan and Visits tabs. Plan and Visits include search, program filters, activity filters, and desktop and mobile navigation.')
h('Scope boundaries')
p('The new Plan dashboard complements the existing weekly, monthly, and quarterly Planner. It summarizes operational activities and networking commitments; it is not a new general task engine. New networking visits do not satisfy operational evidence gates or automatically create workflow assignments.')
p('Contact bulk import, networking Excel export, automated email or calendar invitations, external CRM synchronization, contact deduplication, and contact deletion are not part of the current networking baseline. Production access controls and collaboration behavior need a separate release decision.')

page(3,'Program lifecycle requirements')
p('Programs follow eight ordered stages. Users may inspect stages without changing the actual current stage. Manager completion must run the applicable validations and preserve ownership and evidence requirements.')
table(['Stage','Required outcome before progression'],[
('1  Request and Greenlight','Save drafts separately from submission. Submit a complete request and obtain manager greenlight before curriculum work begins.'),
('2  Curriculum','Record outline, topics, competences, duration, and applicable quality and release evidence.'),
('3  Instructor Setup','Select and prepare instructors, including engagement information, fee approval, contract, and onboarding requirements.'),
('4  Launch Commercial Operations and Finance','Complete launch planning and Go decision, applicable commercial routes, the four operational activities and supervisor review, then finance requirements.'),
('5  Readiness Delivery and Assessment','Confirm readiness; maintain sessions, attendance, and incidents; complete assessment, results, and certificate requirements.'),
('6  Graduate Outcomes','Maintain applicable graduate and recruitment outcomes and supporting evidence.'),
('7  Reporting and Improvement','Complete performance reporting, review, and improvement actions.'),
('8  Final Closure','Confirm operational and financial closure, record lessons, and close the program through manager control.')
],[2.25,4.55])
h('Lifecycle rules')
req('LC01','Controlled handoffs','Stage completion must not bypass required fields, counts, approvals, open incidents, follow-ups, or evidence checks.')
req('LC02','Finance separation','PMB owns actual financial results and final profit and loss closure. Treasury owns collections, settlements, and payment evidence. Both responsibilities remain visible.')
req('LC03','Clear state','Show the actual current stage separately from the selected stage, overall progress, completed-stage count, and individual stage status.')
req('LC04','Relationship records','An operational interaction remains attached to its original program activity. A networking visit is a separate relationship record, even when linked to the same program and activity.')

page(4,'Workspace requirements')
p('Each workspace has a distinct purpose. Navigation must preserve those distinctions and keep related program records reachable without re-entering data.')
table(['Workspace','Required capability'],[
('Manager Dashboard','Summarize portfolio health, attention items, launch decisions, crew work, readiness, delivery, finance, and recent activity.'),
('Programs and Program Workspace','Browse the portfolio and open the controlled lifecycle, requirements, assignments, decisions, and evidence for a program.'),
('My Tasks and Team Members','Show position-owned assignments and personal tasks. Keep team management and member workload views manager-only.'),
('Planner','Provide weekly, monthly, and quarterly task planning within the existing calendar workspace.'),
('Plan','Show in-progress and upcoming operational activities, networking visits, and open visit follow-ups, with filtered summary counts.'),
('Visits','Maintain a reusable contact directory and searchable visit records, including original operational interaction history.'),
('Instructors','Maintain instructor directory and profiles, program engagements, onboarding readiness, and existing manager import tools.'),
('Operations','Expose the four activity lanes, activity status and progress, interactions, and open follow-ups.'),
('Finance and Schedule','Provide program financial summaries and detailed delivery sessions, including dates, times, and statuses.'),
('Documents Reports and History','Find evidence references, review available reports, and inspect recorded workflow history.'),
('Profile Administration and Training','Manage user preferences and themes, configured master data and Position rules, and the existing training walkthrough.')
],[2.15,4.65],9.3)
h('Navigation requirements')
req('NAV01','Desktop and mobile','Expose Plan and Visits in the desktop sidebar and mobile More menu, and identify the active workspace.')
req('NAV02','Default view','Allow Plan or Visits to be selected as the profile default view while retaining existing manager-only navigation restrictions.')
req('NAV03','Contextual actions','Open a visit editor from Plan, open a contact’s visit history from the directory, and open an operational interaction in its original program activity.')

page(5,'Plan dashboard requirements')
p('Plan is an operational planning dashboard. It combines program activities and networking commitments so the user can see what is active, what has not started, and which dated items need attention.')
req('PL01','Dashboard lanes','Provide separate In progress and Upcoming lanes. Include operational activities, networking visits, and open visit follow-ups. Exclude completed and cancelled visit records from these two lanes.')
req('PL02','Activity state','Map an activity’s Not Started status to Upcoming and In Progress to In progress. Show the program, activity name, responsible Position, and program start date. Activity dates are program-level context, not independent deadlines.')
req('PL03','Visit state','Use the visit’s selected status. An overdue Upcoming visit remains Upcoming until edited. A visit dated in the future is not automatically changed to In Progress. Completed visits contribute to the completed summary.')
req('PL04','Follow-up state','Create a planning entry for each visit with a follow-up date that is not completed or cancelled. Dates up to today appear In progress; later dates appear Upcoming. Closing the follow-up removes its open planning entry.')
req('PL05','Summary counts','Show In progress, Upcoming, Overdue, and Completed counts for the filtered records. Overdue means an open item dated before today. A visit and its follow-up are separate planning items; counts are not unique program or contact counts.')
req('PL06','Filters and ordering','Combine program, activity, and case-insensitive text filters. Search titles, contact or owner details, activity names, and program names. Sort open items by date, earliest first, with undated records last. Clear filters restores the full scope.')
req('PL07','Item actions','Use Open activity to navigate to the program’s activity details. Use View visit to open the networking visit editor. Offer Plan a visit and Visit database actions. Require a saved contact before adding a networking visit.')
req('PL08','Empty and long lists','Show an informative empty message when filters match no items. Keep longer desktop lanes scrollable and use a single-column layout on smaller screens.')
h('Primary user journey')
p('Select Plan, narrow the view to a program and activity, inspect the active work and upcoming commitments, and open the relevant record. Update a visit or close its follow-up, save, and see the corresponding counts and lane entries refresh.')
h('Interpretation rule')
p('A completed-visit count reflects recorded visits, including operational history. It does not imply that an operational activity or program stage is complete. Lifecycle completion remains governed by the program workflow.')

page(6,'Visits and contact requirements')
p('Visits is the networking database. One contact can be reused across many visits, programs, and activities, while the visit records preserve the context and outcomes of each interaction.')
req('VI01','Contact directory','Create and edit contact person, organization, role, email, phone, and relationship notes. Require a nonblank contact name and validate an entered email address. Preserve a stable contact identifier when details change.')
req('VI02','Visit creation','Select a saved contact and activity. Choose a program or General networking. Require a visit date and nonblank purpose. Support optional time, location or meeting reference, outcome, additional notes, and follow-up information.')
req('VI03','Visit types and statuses','Support Physical Location, Online Meeting, Phone Call, Event, and Email. Support Upcoming, In Progress, Completed, and Cancelled statuses. Require an outcome before saving a Completed visit.')
req('VI04','Follow-ups','Record a follow-up date, next steps, and completion flag. Reject a follow-up date before the visit date. Retain completed follow-up information in history. Cancelled visits must not contribute open follow-ups to Plan.')
req('VI05','Contact history','Offer History from each contact row to display only that contact’s linked visits, with contact details and notes above the results. Provide an action to return to all contacts’ visits.')
req('VI06','Combined filters','Apply program and activity together. In Visit records, search contact name, organization, email, phone, purpose, outcome, location, notes, activity, and program name. Show visit records by visit date, newest first.')
req('VI07','Directory filtering','Program and activity filters identify contacts through linked visits. Contacts without visits remain visible when those filters are clear. Directory text search includes name, organization, role, email, phone, and relationship notes.')
req('VI08','Existing operational history','Include existing program-activity interactions and label them Operational interaction. Keep their original target text and source record. Open them in the program activity rather than silently converting or duplicating them as networking visits.')
req('VI09','Saving and recovery','Refresh the affected dashboard and directory after editing. Display the existing saving, saved, or save-failed feedback. Persist contact and visit changes across reloads through the application store.')
h('Primary user journey')
p('Add a contact, then add a visit linked to that person and the relevant program and activity. Record the meeting purpose and date. After the visit, enter the outcome, change status, and schedule the next step. Review future or overdue follow-ups in Plan.')

page(7,'Data and validation')
p('Programs provide the program reference and activity context. Contacts and networking visits are stored within masterData.networking in the existing store. Original operational interactions remain within their program activity records.')
table(['Entity','Key fields and rules'],[
('Contact','id; name required; organization; role; email; phone; notes; createdAt; updatedAt when edited. Contact identifiers link visits to the current contact details.'),
('Networking visit','id; contactId required; programId optional; activity required; date required; time optional; status; type; location; objective required; outcome; notes; creation and update metadata.'),
('Visit follow-up','followUp date; followUpNotes; followUpCompleted. No separate contact or visit is created for the follow-up planning entry.'),
('Operational interaction','Original id, target, date, type, objective, outcome, follow-up and evidence, plus the containing program and activity context.'),
('Plan item','Derived view of an activity, visit, or open follow-up. Summary counts and lane entries are calculated from the selected filters; no duplicate Plan database is required.')
],[1.55,5.25])
h('Field constraints')
table(['Field group','Baseline validation'],[
('Contact text','Name, organization, and role: up to 200 characters each. Email: 254. Phone: 60. Notes: 4000.'),
('Visit text','Location or reference: up to 1000 characters. Purpose, outcome, next steps, and additional notes: up to 4000 each.'),
('References','A networking visit must reference an existing contact. A nonempty program reference must identify an available program.'),
('Dates and status','Require visit date; do not allow follow-up before visit date. Require outcome for Completed status.'),
('Display safety','Render user-entered names, notes, and outcomes as text rather than executable markup.')
],[1.55,5.25])
h('Persistence and exports')
p('The baseline uses SQLite as the primary local store with a JSON snapshot. JSON backup export and import include networking contacts and visits. The existing Excel mirror remains focused on personal tasks; networking Excel export is not currently provided.')
p('Contact edits change the displayed contact details across linked networking visits. The baseline does not retain a historical snapshot of a contact’s organization or role at each visit. Historical identity requirements should be decided before a larger shared rollout.')

page(8,'Brand usability and production quality')
p('Retain the original Andalusia Academy bilingual logo without redrawing it, altering its proportions, or substituting another emblem. Use the application’s Terracotta theme as the default visual reference. This Word PRD uses the same original logo and matching colors.')
table(['Theme token','Color','Use'],[
('Primary','#C17A62','Terracotta brand accents'),('Primary dark','#8F4F3B','Primary controls and strong accents'),('Primary soft','#F3E4DF','Light supporting surfaces'),('Cream','#FFF9F6','Warm background'),('Ink','#55372F','Readable brown text')
],[1.7,1.2,3.9])
req('UX01','Responsive layout','Make both new workspaces usable on desktop and phone widths. Contain wide data tables within horizontal scroll regions. Do not introduce horizontal overflow at the page level.')
req('UX02','Accessible controls','Use visible labels, keyboard-focus states, readable status text, keyboard-operable actions, and dialog dismissal with Escape. Do not communicate status by color alone.')
req('UX03','Theme consistency','Use existing theme tokens for surfaces, accents, and borders. Preserve the application’s alternate themes and the original logo assets.')
h('Production requirements beyond the local baseline')
req('NFR01','Identity and authorization','Define authenticated users and enforce access on the server before shared deployment. Existing local profile and Position-based interface restrictions must not be treated as a production authorization boundary.')
req('NFR02','Concurrent editing','Define conflict handling, edit ownership, and reliable retries. The current full-store save model requires review before several users edit shared records at once.')
req('NFR03','Contact data governance','Approve who may view, create, edit, export, archive, or delete contact records; decide retention and audit requirements. The current networking directory is shared locally, and its edits do not provide a complete per-change audit trail.')
req('NFR04','Operational readiness','Agree backup ownership, restore procedures, expected record volumes, and measurable performance targets. Test those agreed volumes before release; no load capacity is claimed by this PRD.')

page(9,'Acceptance criteria and verification')
p('The following criteria define observable behavior for the Plan and Visits enhancement. Baseline checks were exercised with a disposable test database; they do not certify production security, load capacity, or the complete historical lifecycle implementation.')
table(['ID','Scenario and expected result'],[
('AC01','Open Plan and Visits through the desktop sidebar and mobile More menu. The correct workspace becomes active.'),
('AC02','Save a valid contact, edit its phone, and open its history. Details update and only linked visits appear.'),
('AC03','Create a visit for a saved contact, program, and activity. It appears in Visit records and in the appropriate Plan lane.'),
('AC04','Combine program and activity filters. Both conditions apply. An unmatched program gives an empty state; clearing filters restores records.'),
('AC05','Search by contact organization or visit purpose. Matching records remain; unrelated records are excluded.'),
('AC06','Attempt to complete a visit without an outcome or save a follow-up before the visit. The save is blocked and a clear validation message is shown.'),
('AC07','Mark a visit completed with an outcome. Its open visit entry leaves Plan. An unfinished follow-up remains separately visible.'),
('AC08','Complete a visit follow-up. It disappears from open Plan entries while its completed information remains in the visit.'),
('AC09','Load an existing operational interaction. Visits retains its history and identifies its original source; editing follows the existing activity workflow.'),
('AC10','Save and reload the app. Contact and visit data remain. A JSON backup contains the saved networking records.'),
('AC11','Render the tabs at 1440 px and 390 px widths. The root page does not overflow horizontally; the mobile navigation and Escape dialog close work.'),
('AC12','Render a contact containing angle brackets. The value appears as text, and no script or page error is introduced.')
],[.65,6.15],9.3)
h('Evidence and remaining release checks')
p('The local networking test suite passed form, filtering, validation, follow-up, legacy-history, reload, JSON-export, mobile-navigation, and overflow checks. JavaScript syntax checks passed. Desktop and mobile screenshots were visually reviewed.')
p('Before release, perform stakeholder UAT, a complete restore exercise, the agreed permission and concurrency checks, alternate-theme review, and a regression pass for the core lifecycle. Assign owners and record outcomes rather than treating these activities as already approved.')

page(10,'Delivery decisions and review')
h('Release sequence')
req('R1','Local baseline','Deliver Plan and Visits with the current application, retain existing operational interactions, and preserve the Academy brand. This implementation is present in the local workspace.')
req('R2','Stakeholder acceptance','Review the PRD with Academy leadership and relevant Position owners. Confirm the terminology, filtered counts, contact fields, visit states, and separation between networking records and operational evidence.')
req('R3','Shared deployment readiness','Implement or explicitly approve the production requirements for identity, permissions, concurrency, audit history, backups, and hosting. Establish launch and support ownership before a wider rollout.')
h('Decisions to confirm')
table(['Decision','Proposed reviewer'],[
('Should Plan also include personal tasks and independent activity deadlines?','Academy Manager and Operations Supervisor'),
('Who can view, edit, export, archive, or delete shared contact records?','Academy leadership and IT owner'),
('Should legacy operational targets be linked to reusable contacts, and how should duplicates be merged?','Operations and partnerships owners'),
('Should networking follow-ups create assigned tasks or external reminders?','Academy Manager and workflow owners'),
('Are contact import, networking Excel export, and CRM integration required in the next release?','Business owner and IT owner'),
('What record volumes, response targets, retention period, and backup recovery objectives are required?','Business owner and IT owner')
],[4.65,2.15])
h('Source baseline')
p('Prepared from the project README, the program-stage flow revision notes, the application interface and workflow logic in static/index.html, the networking implementation in static/networking.js and static/networking.css, and the store behavior in server.py. Brand assets come from the original static/logo.png and application theme tokens.')
p('Networking verification is recorded by tests/networking.test.cjs and the isolated test server. These sources describe the local project reviewed on 8 September 2026; no external market assumptions or stakeholder sign-offs have been added.')
h('Review outcome')
p('Requested review: confirm the baseline scope, resolve the decisions above, and identify which production requirements are necessary for the intended deployment. Business approval, named delivery owners, and target release dates remain unconfirmed.')

doc.core_properties.title='Andalusia Academy Work Management System Product Requirements'
doc.core_properties.subject='Product scope and requirements including Plan and Visits'
doc.core_properties.author='Andalusia Academy'
doc.core_properties.keywords='Andalusia Academy, WMS, PRD, Plan, Visits'
OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT)
