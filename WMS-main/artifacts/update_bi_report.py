from pathlib import Path
import re, statistics
import pdfplumber
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
ASSETS=ROOT/'artifacts/report-assets'
OUT=ROOT/'artifacts/Andalusia_Academy_Power_BI_Report_Updated.docx'
PDF=Path('C:/Users/eslam.hafny/Downloads/Andalusia_Academy_Power_BI__Report.pdf')
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.left_margin=sec.right_margin=Inches(.65)
sec.top_margin=Inches(1.05);sec.bottom_margin=Inches(.58)
sec.header_distance=Inches(.3);sec.footer_distance=Inches(.25)
NAVY='17364F';TERRA='8F4F3B';INK='24374B';SOFT='F3E4DF'
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 s=doc.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor.from_string('000000')
 for e in list(s.element.iter(qn('w:pBdr'))):e.getparent().remove(e)
normal=doc.styles['Normal'];normal.font.size=Pt(9.5);normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.06
for name,size in [('Title',27),('Subtitle',16),('Heading 1',22),('Heading 2',14),('Heading 3',11)]:
 s=doc.styles[name];s.font.size=Pt(size);s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(7);s.paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(8);doc.styles['Caption'].font.color.rgb=RGBColor.from_string('65758A')
header=sec.header.paragraphs[0];header.paragraph_format.tab_stops.add_tab_stop(Inches(3.9))
header.add_run().add_picture(str(ASSETS/'1_0.png'),height=Inches(.55))
r=header.add_run('\tPower BI and Business Intelligence Handover');r.font.size=Pt(8);r.bold=True;r.font.color.rgb=RGBColor.from_string(NAVY)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Andalusia Academy  |  Version 1.1  |  Confidential - Internal Use  |  Page ');r.font.size=Pt(7.5);r.font.color.rgb=RGBColor.from_string('65758A')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def clean(t):
 t=str(t or '').replace('\u00a0',' ').replace('\ufffd','\u2019')
 t=re.sub(r'^\u2019\s*','',t).replace(' \u2019 ',' - ')
 return t.strip()
def p(t='',style=None,bold=False):
 q=doc.add_paragraph(style=style);q.add_run(t).bold=bold;return q
def h(t):return doc.add_heading(t,2)
def newpage(title=None):
 doc.add_page_break()
 if title:doc.add_heading(title,1)
def para(label,text):
 q=p();q.add_run(label+' ').bold=True;q.add_run(text)
def tab(headers,rows,widths=None,size=8.5):
 widths=widths or [7.2/len(headers)]*len(headers)
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 border=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');border.append(el)
 t._tbl.tblPr.append(border)
 for idx,values in enumerate([headers]+rows):
  row=t.rows[0] if idx==0 else t.add_row();pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if idx==0:pr.append(OxmlElement('w:tblHeader'))
  for c,value,w in zip(row.cells,values,widths):
   c.width=Inches(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),NAVY if idx==0 else ('F7F5F3' if idx%2 else 'FFFFFF'));cp.append(shade)
   mar=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    a=OxmlElement('w:'+side);a.set(qn('w:w'),'65');a.set(qn('w:type'),'dxa');mar.append(a)
   cp.append(mar)
   q=c.paragraphs[0];q.paragraph_format.space_after=Pt(0);q.paragraph_format.line_spacing=1.02
   r=q.add_run(clean(value).replace('\n',' '));r.font.size=Pt(size);r.bold=idx==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if idx==0 else INK)
 p().paragraph_format.space_after=Pt(0)
 return t
def pic(path,caption='',maxh=4.65,width=7.2):
 im=Image.open(path);w,hg=im.size;ww=min(width,maxh*w/hg)
 q=p();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=Pt(3);q.paragraph_format.keep_with_next=bool(caption)
 r=q.add_run().add_picture(str(path),width=Inches(ww));r._inline.docPr.set('descr',caption or Path(path).stem)
 if caption:p(caption,style='Caption')

def normalized_table(t):
 raw=t.extract();cols=[j for j in range(len(raw[0])) if any(row[j] and str(row[j]).strip() for row in raw)]
 if len(cols)<2:return None
 rows=[]
 for rr in raw:
  vals=[clean(rr[j]) for j in cols]
  if not any(vals):continue
  if rows and (not vals[0] or sum(bool(x) for x in vals)==1):
   rows[-1]=[(a+' '+b).strip() for a,b in zip(rows[-1],vals)]
  else:rows.append(vals)
 return rows if len(rows)>1 else None

def original_page(page,number):
 newpage()
 tables=[]
 for t in page.find_tables():
  vals=normalized_table(t)
  if vals:tables.append((t,vals))
 events=[]
 for t,vals in tables:events.append((t.bbox[1],'table',(t,vals)))
 for idx,im in enumerate(page.images):
  if im['top']<75:continue
  files=list(ASSETS.glob(f'{number}_{idx}.*'))
  if files:events.append((im['top'],'image',(im,files[0])))
 lines=page.extract_text_lines()
 groups=[]
 for line in lines:
  if line['top']<78 or line['bottom']>752:continue
  if any(t.bbox[1]-1<=line['top'] and line['bottom']<=t.bbox[3]+1 for t,_ in tables):continue
  text=clean(line['text'])
  if not text or text.startswith('ANDALUSIA ACADEMY |'):continue
  chars=line['chars'];sz=statistics.median(c['size'] for c in chars);bold=sum('Bold' in c['fontname'] for c in chars)>len(chars)*.65
  isbullet=line['text'].lstrip().startswith(('\ufffd','\u2022','\uf0b7'))
  if isbullet:text=text.lstrip('\u2022\uf0b7 ').strip()
  kind='heading' if sz>=13 else ('caption' if text.startswith('Figure ') else 'body')
  if groups and groups[-1]['kind']=='caption' and line['top']-groups[-1]['bottom']<5 and abs(groups[-1]['size']-sz)<1:
   groups[-1]['text']+=' '+text;groups[-1]['bottom']=line['bottom'];continue
  if groups and not isbullet and not groups[-1]['text'].isdigit() and groups[-1]['kind']==kind and abs(groups[-1]['size']-sz)<1 and groups[-1]['bold']==bold and line['top']-groups[-1]['bottom']<(12 if kind=='heading' else 5) and not groups[-1]['text'].startswith('Figure '):
   groups[-1]['text']+=' '+text;groups[-1]['bottom']=line['bottom']
  else:groups.append(dict(text=text,top=line['top'],bottom=line['bottom'],size=sz,bold=bold,kind=kind,bullet=isbullet))
 for g in groups:events.append((g['top'],'text',g))
 for _,kind,v in sorted(events,key=lambda e:e[0]):
  if kind=='table':
   t,rows=v;n=len(rows[0]);widths=([1.6,5.6] if n==2 else ([1.7,2.75,2.75] if n==3 else None))
   if number==41:widths=[1.25,1.85,1.5,1.7,.9]
   if number==2:
    for row in rows:
     if row[0]=='Application baseline':row[1]='Source report: Revision 20.2. Additions: Plan and Visits from the local networking build.'
   tab(rows[0],rows[1:],widths,size=8 if number==41 else 8.5)
  elif kind=='image':
   im,path=v;pic(path,maxh=min(6.9,im['height']/72)-(.55 if number==37 else 0),width=min(7.2,im['width']/72))
  else:
   text=v['text']
   if v['kind']=='heading':
    q=doc.add_heading(text,1 if v['size']>=20 else 2)
   else:
    q=p(text,style='Caption' if v['kind']=='caption' else ('List Bullet' if v['bullet'] else None),bold=v['bold'])
    if not v['bullet'] and v['kind']!='caption':q.runs[0].font.size=Pt(min(10,max(9,v['size'])))
 if number==2:p('Version 1.1 adds sections 3.23 to 3.25, networking data specifications, section 06 KPI Framework, and supplementary BI and mobile mappings.',style='Caption')

def update_overview():
 newpage('Version 1 1 update overview')
 p('Plan and Visits extend the Academy system with an operational planning dashboard and a reusable networking database. This update connects contact people, visit history, outcomes and follow-ups to Programs and activities, and defines how the BI layer should report on them.')
 tab(['Update','Where to find it'],[
 ('Plan dashboard','Section 3.23: In progress and Upcoming lanes, summary cards, search and combined filters.'),
 ('Visits networking database','Section 3.24: visit records, contact directory, contact and visit forms, follow-up behavior and validation.'),
 ('Filter behavior','Section 3.25: Program and activity rules across Plan and Visits.'),
 ('Data and metrics','Section 05 additions and section 06: record fields, model grain, relationships and KPI definitions.'),
 ('Power BI and mobile','Section 07 additions and appendices: analytical pages, acceptance checks and responsive screenshots.')],[1.7,5.5])
 h('Version and access boundary')
 p('The original report documents Revision 20.2, including its authentication and role controls. The new Plan and Visits screens are captured from the local networking build, which uses a local profile and shared networking records. The original account and security descriptions remain the Revision 20.2 baseline; they do not establish that authentication or server-side networking permissions are present in the local build.')
 p('Before these additions are deployed in the Revision 20.2 environment, verify account-based access, ownership and API permissions for contacts, visits and follow-ups. Document the exact deployed application version with the BI refresh configuration.')
 h('Navigation and screenshots')
 p('Plan and Visits are added to the main application navigation and are available through the mobile navigation. Plan opens the Visits database or the visit form. Visits switches between Visit records and Contact directory. The existing weekly, monthly and quarterly Planner remains the task calendar.')
 p('Figures 33 to 39 show actual application screens using demonstration contacts and Programs. They illustrate functionality and are not Academy performance results. Original figures retain their original numbers and version context.')

def new_screens():
 newpage('3 23 Plan dashboard')
 p('Plan consolidates operational activities, visits and open follow-ups into In progress and Upcoming lanes. It supports day-to-day preparation and review across Programs.')
 pic(ASSETS/'new-plan.png','Figure 33 - Plan dashboard with In progress and Upcoming work and Program and activity filters',maxh=4.5)
 para('Summary cards.','In progress and Upcoming count matching work items. Overdue counts open dated items before today. Completed counts completed activities and visits. A visit and its open follow-up are separate planning items.')
 para('Work cards.','Each card shows its kind, status, title, Program, activity, contact or owner detail, and date. Cards sort by date with undated items last. Open activity navigates to the operational activity; View visit opens the networking visit form.')
 para('Quick actions.','Visit database opens Visits. Plan a visit opens the visit form and carries across the selected Program and activity. If no contact exists, the contact form opens first.')
 newpage('3 23 Plan behavior and status rules')
 tab(['Source','Plan behavior'],[
 ('Operational activity','Not Started maps to Upcoming; In Progress remains In progress. The displayed activity date is the Program start date, not a separately maintained activity due date.'),
 ('Networking visit','Uses the visit status and date entered by the user: Upcoming, In Progress, Completed or Cancelled. Dates do not automatically change the visit status.'),
 ('Operational visit history','Existing operational interactions are displayed as Completed visits and open in their original activity.'),
 ('Open follow-up','Created as a derived planning item when a follow-up date exists, completion is false and the visit is not Cancelled. Due today or earlier maps to In progress; a future date maps to Upcoming.'),
 ('Completed follow-up','Stops appearing as an open Plan item. It is not a separate entry in the Completed card.'),
 ('Overdue','An additional flag for Upcoming or In Progress items whose date is before today. It overlaps the two open lanes; it is not a third exclusive status.')],[1.5,5.7])
 h('Search and filters')
 p('Search matches planning titles, displayed contact or owner details, activity and Program name. Program, activity and search conditions apply together. Clear filters restores the complete view. Empty lanes show a clear no-results message.')
 h('Operational boundaries')
 p('Plan is a view of existing operational and networking records. It does not create a separate planning database or include personal tasks from the task Planner. Networking follow-ups do not create workflow assignments and do not satisfy operational evidence or stage-completion requirements. Operational follow-ups retain their existing workflow behavior.')
 h('BI interpretation')
 p('Use separate measures for activities, visits and follow-ups, then a clearly labelled combined work-item count. Retain source type and record keys so a visit and its follow-up cannot be mistaken for duplicate visits. Refresh the date comparison consistently using the approved business timezone.')
 newpage('3 24 Visits networking database')
 p('Visits stores interactions and their next steps alongside reusable contact people. Users can review relationships across Programs or maintain general networking visits without a Program.')
 pic(ASSETS/'new-visits.png','Figure 34 - Visits database with visit records, outcomes and open follow-ups',maxh=4.5)
 para('Visit records.','The table displays date and time, status, contact person and organization, Program and activity, purpose, visit type, location, outcome, notes, follow-up and available action. Records sort from newest visit date to oldest.')
 para('Actions.','Add contact creates a person. Add visit schedules or records an interaction. Edit visit changes a networking record. Open activity returns an existing operational interaction to its original workflow.')
 para('Summary cards.','Contacts counts the matching reusable directory records. Visits counts visible visit records including operational history. Upcoming counts visible visits with that status. Open follow-ups counts dated, unfinished follow-ups on visits that are not Cancelled.')
 para('History.','Selecting History in the contact directory restricts visit records to that contact and shows contact details above the table. Show all contacts’ visits removes that selection.')
 newpage('3 24 Contact directory')
 pic(ASSETS/'new-contacts.png','Figure 35 - Contact directory with contact details, visit counts and history actions',maxh=4.5)
 para('Reusable identity.','A contact has a stable ID and can be linked to multiple visits across Programs and activities. Editing contact details updates their display on linked visits without creating a new person.')
 para('Directory actions.','History shows that person’s visit records. Edit opens the contact form. Add visit opens the visit form with the person already selected. The directory visit-count column shows the person’s total linked visits, even when Program or activity filters narrow the directory.')
 para('Filtering.','Program and activity filters select contacts through their linked visits. Contacts without visits appear when these filters are cleared. Directory search matches name, organization, role, email, phone and notes.')
 para('Current scope.','Contact deletion, merging duplicates, bulk import and automatic deduplication are not implemented in this addition. Operational visit targets are not automatically converted into reusable contacts.')
 newpage('3 24 Contact form')
 pic(ASSETS/'new-contact-form.png','Figure 36 - Contact form for creating or editing a reusable contact person',maxh=4.5)
 tab(['Field','Requirement and behavior'],[
 ('Contact name','Required. Up to 200 characters; whitespace is trimmed.'),
 ('Organization and role','Optional. Up to 200 characters each.'),
 ('Email and phone','Optional. Email uses the email input validation and a 254-character limit; phone allows up to 60 characters.'),
 ('Notes','Optional relationship context, up to 4,000 characters.'),
 ('Save contact','Creates a stable ID and creation timestamp for a new person, or updates the existing record and update timestamp.'),
 ('Cancel or close','Closes the form without saving the entered changes.')],[1.6,5.6])
 newpage('3 24 Visit form and follow ups')
 pic(ASSETS/'new-visit-form.png','Figure 37 - Visit form with Program, activity, status, purpose, outcome and follow-up controls',maxh=4.55)
 para('Required inputs.','Select an existing contact and an activity, then enter visit date and purpose. Program is optional and defaults to General networking. If selected, the Program must still exist.')
 para('Visit details.','Time, location or meeting reference and additional notes are optional. Status is Upcoming, In Progress, Completed or Cancelled. Visit type is Physical Location, Online Meeting, Phone Call, Event or Email. A Completed visit requires an outcome.')
 para('Follow-up.','Enter an optional follow-up date and next steps. The follow-up date must be on or after the visit date. Mark Follow-up completed when the action is finished. A Completed visit can still have an open follow-up; Cancelled visits are excluded from open follow-up counts.')
 para('Save.','Valid records refresh Plan and Visits and are persisted with the application store. Missing required inputs or invalid date relationships keep the form open for correction. Updating a visit retains its ID.')
 newpage('3 25 Program and activity filters')
 p('Program and activity filtering is available in both new tabs. Selections narrow the displayed data and the associated summaries; users do not need separate databases for each Program.')
 tab(['Control or context','Expected behavior'],[
 ('Program','All programs removes the Program restriction. Selecting a Program filters by its stable ID.'),
 ('Activity','All activities removes the activity restriction. Options include standard activities and available activity names from the selected Program and recorded visits.'),
 ('Combined selection','Program AND activity AND search must all match. A matching activity under another Program is excluded.'),
 ('Change Program','Rebuilds the available activity options. A selected activity is retained only while it is still an available option.'),
 ('General networking','Visits without a Program appear under All programs and can still be filtered by activity.'),
 ('Contact history','Further restricts the visible visit records to one contact. Clear filters also clears the selected contact history.'),
 ('Clear filters','Removes text, Program and activity restrictions and refreshes the active tab.'),
 ('Empty results','Shows an explicit no-results message and provides the existing clear-filter and add-record actions.')],[1.65,5.55])
 h('Search and count differences')
 p('Visit search matches contact name, organization, email, phone, purpose, outcome, location, notes, activity and Program name. Directory search includes the contact’s role and notes. The Contacts card follows the directory filter logic; selecting one person’s History does not change it into a count of the selected person. The application-wide search is separate and does not include networking contacts or visits.')
 h('Example')
 p('Selecting Clinical Leadership and Instructors returns only instructor-related items for that Program. A Sponsorships visit in the same Program is excluded. Changing to All activities restores it. A contact linked only to another Program is excluded from the filtered directory.')

def data_additions():
 newpage('05 Networking data additions')
 p('The networking extension is persisted under masterData.networking in the application configuration and JSON store. It contains contacts and visits arrays. Existing operational interactions remain inside each Program’s activity records.')
 tab(['Entity and grain','Stored fields'],[
 ('Contact\nOne row per contact ID','id; name; organization; role; email; phone; notes; createdAt; updatedAt when edited.'),
 ('Networking visit\nOne row per visit ID','id; contactId; programId or blank; activity; date; time; status; type; location; objective; outcome; notes; followUp; followUpNotes; followUpCompleted; createdAt; createdBy; updatedAt when edited.'),
 ('Operational interaction\nOne row per source visit in a Program activity','Existing visit fields including ID, target, date, type, objective, outcome, evidence and follow-up fields. Program and activity context come from the containing record.'),
 ('Plan item\nOne derived row per activity, visit or open follow-up','Source kind and key; Program; activity; title; detail; effective status; display date. This is a derived view, not a stored Plan table.')],[2,5.2])
 h('Relationships')
 p('A contact has many networking visits through contactId. A Program has many linked networking visits through programId, while General networking visits have no Program. Operational interactions may have a free-text target without contactId; retain them as unmatched contacts until an explicit mapping is approved.')
 h('Data quality and privacy')
 p('Validate unique contact and visit IDs, existing contact references, valid optional Program references, permitted statuses, required purpose/date/activity, outcomes for completed visits and valid follow-up dates. Do not infer a unique person from name alone. Limit contact email, phone and free-text notes to approved BI audiences.')
 newpage('05 Networking analytical model')
 tab(['Model object','Grain and relationship rule'],[
 ('DimContact','One row per contact ID. Link to networking visits. Use an explicit unknown/unmatched member for operational targets without a mapped contact.'),
 ('DimProgram and DimActivity','Reuse the Program dimension. Create a controlled activity mapping; retain raw activity labels. Use an explicit General networking member for blank Program links.'),
 ('FactVisits','One row per source visit. Use a composite key including source type, Program/activity context where needed, and source ID to avoid ID collisions.'),
 ('FactFollowUps','At most one current follow-up per source visit in this implementation. Store follow-up date, completed flag and source visit key; exclude Cancelled visits from the open measure.'),
 ('FactPlanItems or analytical view','Union activity rows, visit rows and open follow-up rows with distinct item-type keys. Report type-specific totals alongside any combined count.'),
 ('DimDate','Use visit date for visits and follow-up date for follow-up measures. Use role-playing date relationships or explicit measure logic to avoid ambiguous date slicing.'),
 ('History snapshots','Add dated snapshots or change events for time-series status, overdue trends and changing contact attributes. Current-state arrays alone do not provide full history.')],[1.65,5.55])
 h('Refresh and reconciliation')
 p('Extract contacts and visits from the JSON configuration, flatten operational activity visits separately, then apply a shared Program and activity mapping. Preserve source type throughout transformation. Reconcile visit counts before adding follow-up rows, then validate each KPI against the application using the same filters and reference date.')
 h('Availability limits')
 p('Networking follow-ups have a completed flag but no dedicated completion timestamp. Do not publish on-time follow-up completion, historical conversion or interaction-to-revenue attribution as measured facts without adding the required event dates and business links. Do not include password hashes, salts or session secrets in the BI model.')

def kpis():
 newpage('06 KPI Framework')
 p('Each measure must declare its population, filter context, time basis and source. Current application cards are snapshots. Historical trends require dated facts, snapshots or reliable event timestamps.')
 tab(['Measure','Calculation and interpretation'],[
 ('Plan In progress','Count filtered Plan rows whose effective status is In Progress, including activities, visits and open follow-ups due today or earlier.'),
 ('Plan Upcoming','Count filtered Plan rows whose effective status is Upcoming. Includes Not Started activities, Upcoming visits and future open follow-ups.'),
 ('Plan Overdue','Count filtered open Plan rows with a nonblank display date before the reference date. This is a subset of In progress plus Upcoming; do not add it to those totals.'),
 ('Plan Completed','Count filtered completed activities and visits. Completed follow-ups are not retained as separate completed Plan rows.'),
 ('Directory contacts','Distinct stored contact IDs matching directory search and linked-visit Program/activity scope. Unvisited contacts are included only when Program/activity restrictions are cleared.'),
 ('Visible visits','Count filtered networking and operational visit rows, after any selected contact-history restriction. Follow-up rows do not increase this measure.'),
 ('Upcoming visits','Count visible visits with status Upcoming. Do not substitute a future-date test for the explicit status.'),
 ('Open follow-ups','Count visible visits with a follow-up date, completed flag false and status other than Cancelled.'),
 ('Overdue follow-ups','Proposed BI measure: open follow-ups with a due date before the agreed reference date.'),
 ('Contacts with visits','Proposed BI measure: distinct contact IDs on filtered networking visits. Exclude unmatched operational targets unless mapped.')],[1.65,5.55])
 h('Reference date')
 p('Approve the reporting timezone and date boundary with the Academy. The local UI uses its current application date. A BI refresh must record its as-of date and apply it consistently, especially for follow-ups due today and overdue comparisons.')
 newpage('06 Portfolio KPI definitions and safeguards')
 tab(['KPI family','Definition or prerequisite'],[
 ('Program completion','Average Program progress across the selected portfolio. Preserve the eight-stage model and agree whether the business wants equal stage weights or another approved weighting.'),
 ('Task completion','Completed task count divided by total tasks in the defined population. Do not double-count assignments and their mirrored personal tasks.'),
 ('Team utilization','The local dashboard uses a workload proxy based on open work. A time-capacity utilization KPI requires available hours and assigned/planned hours; label the proxy clearly.'),
 ('Delivery this month','Count sessions scheduled in the selected month. Delivered sessions must use completion status and a separately labelled measure.'),
 ('Budget utilization','Actual cost divided by budget for the defined Program population. Return blank or a clearly stated exception when budget is zero.'),
 ('Net result and margin','Local financial basis: revenue plus sponsorship minus cost. Margin divides net result by revenue plus sponsorship. Verify that sponsorship is not already included in revenue before loading consolidated financial facts.'),
 ('Stage aging and SLA','Require stage entry/exit dates and an approved SLA calendar. Current progress or selected stage cannot substitute for elapsed duration.'),
 ('On-time completion','Requires original or governed due dates and completion timestamps. A current overdue count does not prove historical on-time performance.'),
 ('Targets and comparisons','Business owners must approve targets, denominators, period filters and exclusions. Fixed dashboard benchmarks are not evidence of approved targets or historical performance.')],[1.75,5.45])
 h('Measure acceptance')
 p('Check empty populations, zero denominators, missing dates, Cancelled visits, unlinked contacts and Programs, completed visits with open follow-ups, and overlapping overdue counts. Compare Program-only, activity-only and combined-filter totals. Display the refresh timestamp and explain any unavailable measure instead of estimating unsupported results.')

def bi_additions():
 newpage('07 Planning and networking BI requirements')
 tab(['BI page','Required analytical experience'],[
 ('Plan and Activity Overview','Separate In progress, Upcoming, Overdue and Completed cards; activity/visit/follow-up breakdown; dated work list; Program and activity slicers; open-item aging where a valid date exists.'),
 ('Visits and Networking','Visit status and type, Program/activity mix, recent interactions, upcoming visits, open/overdue follow-ups, contact and organization coverage.'),
 ('Contact Detail','Contact profile within authorized scope; chronological visit history; linked Programs and activities; outcomes; next steps and unresolved follow-ups.'),
 ('Operations Performance','Retain the original operational lane metrics. Clearly distinguish operational evidence visits from standalone networking visits.'),
 ('Data Quality','Missing or invalid references, duplicate candidates, missing completed-visit outcomes, unavailable dates, unmapped activity labels and unmatched operational targets.')],[1.8,5.4])
 h('Interactions and filtering')
 p('Provide Program, activity, visit type, visit status and relevant date slicers. Contact and organization selections should filter networking facts through approved relationships. Support drill-through from a Program or contact into its visit records and follow-ups, with clear reset-filter controls and an as-of date.')
 h('Security and delivery')
 p('Apply the approved account and organizational scope to datasets and exports. Confirm whether contact details and notes belong in executive reports. Verify the new networking endpoints against the authentication-enabled target release before treating them as a secured production source.')
 h('Acceptance checks')
 p('Reconcile app and BI counts for the same Program/activity/search scope and reference date; test a General networking visit; a contact with no visits; a Completed visit with an open follow-up; and a Cancelled visit with a follow-up date. Confirm that visit totals are unchanged when the model adds follow-up facts and that operational history is counted once.')

def mobile_additions():
 newpage('A Plan and Visits mobile screens')
 q=p();q.alignment=WD_ALIGN_PARAGRAPH.CENTER
 for name in ['new-plan-mobile','new-visits-mobile']:
  r=q.add_run();r.add_picture(str(ASSETS/(name+'.png')),height=Inches(6.4));q.add_run('   ')
 p('Figures 38 and 39 - Plan and Visits at a 390-pixel mobile viewport, using demonstration records',style='Caption')
 p('Plan stacks the summary cards and planning lanes. Visits retains the contact and visit controls with an internally scrollable records table. Program and activity filters remain available. These additions were captured at 1440-pixel desktop and 390-pixel mobile widths.')
def mapping_additions():
 newpage('B Additional UI to BI mapping')
 tab(['Application screen','Data and model','BI destination'],[
 ('Plan dashboard','Derived activity, visit and follow-up rows; Program, activity and date dimensions','Plan and Activity Overview'),
 ('Visits records','FactVisits, FactFollowUps, DimContact, DimProgram, DimActivity and visit/follow-up date roles','Visits and Networking'),
 ('Contact directory','DimContact and linked visits; organization and relationship role attributes','Networking coverage and Contact Detail'),
 ('Contact form','Contact identity and current attributes; created/updated metadata where present','Contact data quality'),
 ('Visit form','Visit status, purpose, outcome, type, dates, follow-up fields and Program/activity context','Visit performance and follow-up exceptions'),
 ('Contact history','Selected contact and linked chronological interactions','Contact Detail drill-through'),
 ('Program and activity filters','Shared conformed dimensions and explicit AND logic','All planning and networking analytical pages')],[1.5,3.25,2.45])
 h('Release verification checklist')
 for text in [
 'Create and edit a contact; confirm linked visits display the updated contact details.',
 'Save an Upcoming visit; complete it only after entering an outcome; verify follow-up date validation.',
 'Complete a follow-up and confirm it leaves the open Plan lanes and open follow-up count.',
 'Apply Program and activity filters together; clear filters; verify empty states and contact-history reset.',
 'Reload the application and restore a JSON backup in a test environment; confirm contacts and visits remain linked.',
 'Verify desktop and mobile navigation, form access and internal table scrolling.',
 'Verify authentication and authorization in the exact target release, then reconcile BI measures against that release.'
 ]:p(text,style='List Bullet')

# Cover retains the Academy artwork from the supplied PDF.
pic(ASSETS/'1_1.png',maxh=2.2,width=2.2)
p('ANDALUSIA ACADEMY',bold=True).alignment=WD_ALIGN_PARAGRAPH.CENTER
q=p('Andalusia Management System',style='Title');q.alignment=WD_ALIGN_PARAGRAPH.CENTER
q=p('Power BI and Business Intelligence Handover Report',style='Subtitle');q.alignment=WD_ALIGN_PARAGRAPH.CENTER
p('Operational System, UI Catalogue, Data Model, KPI Framework and BI Requirements').alignment=WD_ALIGN_PARAGRAPH.CENTER
p()
tab(['Document control','Detail'],[('Version','1.1'),('Date','13 September 2026'),('Original edition','23 August 2026 - Andalusia Management System (AMS)'),('Prepared for','Power BI / Business Intelligence Team'),('Update','Plan dashboard, Visits networking database and new screen captures')],[2,5.2],size=10)
p('The report connects operational functionality with analytical requirements. This edition adds planning and networking workflows, their data definitions and reporting measures, while retaining the original Academy system catalogue and branding.')
with pdfplumber.open(PDF) as source:
 for number,page in enumerate(source.pages,1):
  if number==1:continue
  original_page(page,number)
  if number==2:update_overview()
  if number==31:new_screens()
  if number==35:data_additions();kpis()
  if number==38:bi_additions()
  if number==40:mobile_additions()
mapping_additions()
doc.core_properties.title='Andalusia Academy Power BI and Business Intelligence Handover Report'
doc.core_properties.subject='Version 1.1 with Plan and Visits networking updates'
doc.core_properties.author='Andalusia Academy'
doc.core_properties.version='1.1'
doc.save(OUT)
print(OUT)
print('Paragraphs',len(doc.paragraphs),'Tables',len(doc.tables),'Images',len(doc.inline_shapes))
